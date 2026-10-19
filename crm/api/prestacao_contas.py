# Copyright (c) 2026, Stratcompany and contributors
# For license information, please see license.txt
#
# Prestação de contas ao cliente: o que o escritório recebeu em nome do cliente (alvará, acordo,
# RPV, precatório), quanto ficou de honorários, quais custas adiantadas foram descontadas e quanto
# foi (ou falta ser) repassado. Gera o demonstrativo em PDF para entregar ao cliente e o
# relatório mensal do carnê-leão (honorários recebidos de pessoa física).

import base64
import io
import json

import frappe
from frappe import _
from frappe.utils import add_months, cint, flt, get_first_day, get_last_day, getdate, nowdate

MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]


def _managers_only():
	frappe.only_for(("System Manager", "Sales Manager"))


def _brl(v) -> str:
	return "R$ " + f"{flt(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _data_br(d) -> str:
	return getdate(d).strftime("%d/%m/%Y") if d else ""


def _cliente(deal: str) -> str:
	from crm.api.automacoes import _client

	return _client(deal).completo or deal


# ------------------------------------------------------------------ leitura

@frappe.whitelist()
def resumo(deal: str) -> dict:
	_managers_only()
	doc = frappe.get_doc("CRM Deal", deal)
	recebimentos = []
	for r in frappe.get_all("CRM Recebimento", filters={"deal": deal}, pluck="name", order_by="data desc, creation desc"):
		d = frappe.get_doc("CRM Recebimento", r)
		recebimentos.append({
			"name": d.name, "data": str(d.data), "origem": d.origem, "valor_bruto": flt(d.valor_bruto),
			"percentual_honorarios": flt(d.percentual_honorarios), "valor_honorarios": flt(d.valor_honorarios),
			"despesas_reembolsadas": flt(d.despesas_reembolsadas), "valor_repasse": flt(d.valor_repasse),
			"valor_parceiros": flt(d.valor_parceiros), "repasse_status": d.repasse_status,
			"repasse_data": str(d.repasse_data or ""), "repasse_forma": d.repasse_forma or "", "observacoes": d.observacoes or "",
			"parceiros": [{"parceiro": p.parceiro, "oab": p.oab or "", "percentual": flt(p.percentual), "valor": flt(p.valor),
						   "pago": cint(p.pago)} for p in d.parceiros],
		})
	despesas = frappe.get_all(
		"CRM Despesa", filters={"deal": deal},
		fields=["name", "descricao", "categoria", "valor", "status", "data_vencimento", "reembolsavel", "reembolsado", "recebimento"],
		order_by="data_vencimento desc",
	)
	pendentes = [d for d in despesas if cint(d.reembolsavel) and not d.recebimento]
	return {
		"cliente": _cliente(deal),
		"percentual_exito": flt(doc.get("percentual_exito")),
		"recebimentos": recebimentos,
		"despesas": [{**d, "valor": flt(d.valor), "data_vencimento": str(d.data_vencimento or "")} for d in despesas],
		"totais": {
			"recebido": sum(r["valor_bruto"] for r in recebimentos),
			"honorarios": sum(r["valor_honorarios"] for r in recebimentos),
			"despesas_descontadas": sum(r["despesas_reembolsadas"] for r in recebimentos),
			"repassado": sum(r["valor_repasse"] for r in recebimentos if r["repasse_status"] == "Repassado"),
			"a_repassar": sum(r["valor_repasse"] for r in recebimentos if r["repasse_status"] != "Repassado"),
			"custas_a_descontar": sum(flt(d.valor) for d in pendentes),
		},
	}


# ------------------------------------------------------------------ gravação

@frappe.whitelist()
def salvar_percentual(deal: str, percentual) -> dict:
	_managers_only()
	pct = flt(percentual)
	if pct < 0 or pct > 100:
		frappe.throw(_("O percentual precisa estar entre 0 e 100."))
	frappe.db.set_value("CRM Deal", deal, "percentual_exito", pct)
	return {"ok": True}


@frappe.whitelist()
def salvar_recebimento(dados) -> dict:
	"""Cria ou atualiza um recebimento. A conta (honorários, custas, repasse) é feita no registro."""
	_managers_only()
	dados = frappe.parse_json(dados) or {}
	campos = ("data", "origem", "valor_bruto", "percentual_honorarios", "valor_honorarios", "repasse_status",
			  "repasse_data", "repasse_forma", "observacoes")
	if dados.get("name"):
		doc = frappe.get_doc("CRM Recebimento", dados["name"])
	else:
		doc = frappe.new_doc("CRM Recebimento")
		doc.deal = dados.get("deal")
	for c in campos:
		if c in dados:
			doc.set(c, dados.get(c) or None)
	if "percentual_honorarios" in dados and "valor_honorarios" not in dados:
		doc.valor_honorarios = 0  # recalcula pelo percentual
	if "parceiros" in dados:
		doc.set("parceiros", [])
		for p in dados.get("parceiros") or []:
			if (p.get("parceiro") or "").strip():
				doc.append("parceiros", {
					"parceiro": p["parceiro"].strip()[:140], "oab": (p.get("oab") or "").strip()[:40],
					"percentual": flt(p.get("percentual")), "valor": flt(p.get("valor")) if not flt(p.get("percentual")) else 0,
					"pago": cint(p.get("pago")),
				})
	doc.save(ignore_permissions=True)
	return {"name": doc.name, "valor_repasse": flt(doc.valor_repasse), "valor_honorarios": flt(doc.valor_honorarios)}


@frappe.whitelist()
def marcar_repasse(recebimento: str, data: str | None = None, forma: str | None = None) -> dict:
	_managers_only()
	doc = frappe.get_doc("CRM Recebimento", recebimento)
	doc.repasse_status = "Repassado"
	doc.repasse_data = data or nowdate()
	if forma:
		doc.repasse_forma = forma
	doc.save(ignore_permissions=True)
	return {"ok": True}


@frappe.whitelist()
def excluir_recebimento(recebimento: str) -> dict:
	_managers_only()
	frappe.delete_doc("CRM Recebimento", recebimento, ignore_permissions=True)
	return {"ok": True}


@frappe.whitelist()
def nova_custa(deal: str, descricao: str, valor, data: str | None = None, categoria: str | None = None, reembolsavel=1) -> dict:
	"""Custa ou despesa paga pelo escritório para este caso (por padrão, para cobrar do cliente)."""
	_managers_only()
	if flt(valor) <= 0 or not (descricao or "").strip():
		frappe.throw(_("Informe a descrição e o valor."))
	opcoes = (frappe.get_meta("CRM Despesa").get_field("categoria").options or "").split("\n")
	cat = categoria if categoria in opcoes else ("Custas e Taxas" if "Custas e Taxas" in opcoes else "Outros")
	d = frappe.get_doc({
		"doctype": "CRM Despesa", "deal": deal, "descricao": descricao.strip()[:140], "valor": flt(valor),
		"categoria": cat, "status": "Pago", "data_vencimento": data or nowdate(), "data_pagamento": data or nowdate(),
		"reembolsavel": cint(reembolsavel),
	}).insert(ignore_permissions=True)
	return {"name": d.name}


@frappe.whitelist()
def marcar_reembolsavel(despesa: str, reembolsavel) -> dict:
	_managers_only()
	if frappe.db.get_value("CRM Despesa", despesa, "recebimento"):
		frappe.throw(_("Essa custa já foi descontada numa prestação de contas."))
	frappe.db.set_value("CRM Despesa", despesa, "reembolsavel", cint(reembolsavel))
	return {"ok": True}


# ------------------------------------------------------------------ demonstrativo (PDF)

def _markdown_demonstrativo(deal: str, recebimento: str | None = None) -> str:
	r = resumo(deal)
	itens = [x for x in r["recebimentos"] if not recebimento or x["name"] == recebimento]
	if not itens:
		frappe.throw(_("Ainda não há recebimento registrado neste caso."))
	from crm.api.modelos_documentos import _client_data

	dados = _client_data(frappe.get_doc("CRM Deal", deal))
	linhas = [
		"# PRESTAÇÃO DE CONTAS",
		"",
		f"**Cliente:** {r['cliente']}" + (f" — CPF/CNPJ {dados['cpf_cnpj']}" if dados.get("cpf_cnpj") else ""),
		"",
		f"**Data:** {_data_br(nowdate())}",
		"",
	]
	for i, x in enumerate(itens, 1):
		linhas += [
			f"## {i}. {x['origem']} recebido em {_data_br(x['data'])}",
			"",
			f"- Valor recebido: **{_brl(x['valor_bruto'])}**",
			f"- Honorários advocatícios ({x['percentual_honorarios']:g}%): {_brl(x['valor_honorarios'])}" if x["percentual_honorarios"]
			else f"- Honorários advocatícios: {_brl(x['valor_honorarios'])}",
			f"- Custas e despesas adiantadas pelo escritório: {_brl(x['despesas_reembolsadas'])}",
			f"- **Valor líquido do cliente: {_brl(x['valor_repasse'])}**",
			f"- Repasse: {'realizado em ' + _data_br(x['repasse_data']) + (' via ' + x['repasse_forma'] if x['repasse_forma'] else '') if x['repasse_status'] == 'Repassado' else 'a realizar'}",
			"",
		]
		custas = [d for d in r["despesas"] if d.get("recebimento") == x["name"]]
		if custas:
			linhas += ["Custas descontadas:", ""] + [f"- {_data_br(d['data_vencimento'])} — {d['descricao']}: {_brl(d['valor'])}" for d in custas] + [""]
	if len(itens) > 1:
		linhas += [
			"## Resumo",
			"",
			f"- Total recebido: {_brl(sum(x['valor_bruto'] for x in itens))}",
			f"- Total de honorários: {_brl(sum(x['valor_honorarios'] for x in itens))}",
			f"- Total de custas descontadas: {_brl(sum(x['despesas_reembolsadas'] for x in itens))}",
			f"- **Total do cliente: {_brl(sum(x['valor_repasse'] for x in itens))}**",
			"",
		]
	linhas += ["", "Declaro ter recebido os valores acima e estar de acordo com esta prestação de contas.", "",
			   "", "______________________________", r["cliente"]]
	return "\n".join(linhas)


@frappe.whitelist()
def demonstrativo_pdf(deal: str, recebimento: str | None = None) -> dict:
	_managers_only()
	from crm.api.modelos_documentos import _build

	pdf = _build(_markdown_demonstrativo(deal, recebimento), {}, "Outro", "Prestação de contas", "", False)
	nome = f"Prestação de contas - {_cliente(deal)}.pdf"
	return {"pdf": base64.b64encode(pdf).decode(), "nome": nome}


# ------------------------------------------------------------------ carnê-leão

def _pessoa_fisica(deal) -> bool:
	if deal.get("organization"):
		return False
	from crm.api.modelos_documentos import _client_data

	doc = "".join(c for c in _client_data(deal).get("cpf_cnpj", "") if c.isdigit())
	return len(doc) != 14  # 14 dígitos = CNPJ


@frappe.whitelist()
def carne_leao(mes: str | None = None) -> dict:
	"""Honorários recebidos de pessoa física no mês (base do carnê-leão). mes = 'AAAA-MM'."""
	_managers_only()
	ref = getdate(f"{mes}-01") if mes else get_first_day(nowdate())
	ini, fim = get_first_day(ref), get_last_day(ref)
	hons = frappe.get_all(
		"CRM Honorario",
		filters={"status": "Pago", "data_pagamento": ["between", [ini, fim]]},
		fields=["name", "deal", "valor", "data_pagamento", "tipo_honorario", "observacoes"],
		order_by="data_pagamento asc",
	)
	from crm.api.modelos_documentos import _client_data

	linhas, pj = [], 0.0
	cache = {}
	for h in hons:
		if h.deal not in cache:
			cache[h.deal] = frappe.get_doc("CRM Deal", h.deal) if frappe.db.exists("CRM Deal", h.deal) else None
		deal = cache[h.deal]
		if not deal:
			continue
		if not _pessoa_fisica(deal):
			pj += flt(h.valor)
			continue
		linhas.append({
			"data": str(h.data_pagamento), "cliente": _cliente(deal.name), "cpf": _client_data(deal).get("cpf_cnpj", ""),
			"valor": flt(h.valor), "tipo": h.tipo_honorario, "referencia": h.observacoes or "",
		})
	return {
		"mes": f"{MESES[ini.month - 1]} de {ini.year}",
		"inicio": str(ini),
		"linhas": linhas,
		"total_pf": round(sum(x["valor"] for x in linhas), 2),
		"total_pj": round(pj, 2),
		"sem_cpf": sum(1 for x in linhas if not x["cpf"]),
	}


@frappe.whitelist()
def carne_leao_planilha(mes: str | None = None) -> dict:
	r = carne_leao(mes)
	from openpyxl import Workbook
	from openpyxl.styles import Font

	wb = Workbook()
	ws = wb.active
	ws.title = "Carnê-leão"
	ws.append([f"Honorários recebidos de pessoa física — {r['mes']}"])
	ws["A1"].font = Font(bold=True)
	ws.append([])
	ws.append(["Data", "Cliente", "CPF", "Valor", "Tipo", "Referência"])
	for c in ws[3]:
		c.font = Font(bold=True)
	for x in r["linhas"]:
		ws.append([_data_br(x["data"]), x["cliente"], x["cpf"], x["valor"], x["tipo"], x["referencia"]])
	ws.append([])
	ws.append(["", "Total de pessoa física", "", r["total_pf"]])
	ws.append(["", "Recebido de pessoa jurídica (fora do carnê-leão)", "", r["total_pj"]])
	for col, w in zip("ABCDEF", (12, 34, 18, 14, 18, 40)):
		ws.column_dimensions[col].width = w
	buf = io.BytesIO()
	wb.save(buf)
	return {"arquivo": base64.b64encode(buf.getvalue()).decode(), "nome": f"Carnê-leão - {r['mes']}.xlsx"}

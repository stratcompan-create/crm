# Copyright (c) 2026, Stratcompany and contributors
# For license information, please see license.txt
#
# Segurança e dados do escritório:
# - Login em duas etapas (código no aplicativo autenticador do celular).
# - Exportar todos os dados do escritório numa planilha ("os dados são seus").
# - LGPD: consentimento registrado, quem abriu cada caso, dados de um titular a pedido e
#   anonimização (o financeiro fica, porque a lei obriga a guardar).
# - Carga de trabalho de cada pessoa da equipe.

import base64
import io
import json

import frappe
from frappe import _
from frappe.utils import add_days, cint, flt, get_first_day, now_datetime, nowdate

PAPEIS_2FA = ("System Manager", "Sales Manager", "Sales User")


def _gestor():
	frappe.only_for(("System Manager", "Sales Manager"))


def carimbar_consentimento(doc, method=None):
	"""Guarda quando a pessoa autorizou o uso dos dados."""
	if doc.meta.has_field("consentimento_lgpd"):
		if cint(doc.get("consentimento_lgpd")) and not doc.get("consentimento_em"):
			doc.consentimento_em = now_datetime()
		elif not cint(doc.get("consentimento_lgpd")):
			doc.consentimento_em = None


# ------------------------------------------------------------------ login em duas etapas

@frappe.whitelist()
def estado() -> dict:
	_gestor()
	return {
		"dois_fatores": cint(frappe.db.get_single_value("System Settings", "enable_two_factor_auth")),
		"demonstracao": bool(frappe.conf.get("demo_username")),
		"casos_sigilosos": frappe.db.count("CRM Deal", {"sigiloso": 1}) if frappe.get_meta("CRM Deal").has_field("sigiloso") else 0,
	}


@frappe.whitelist()
def definir_dois_fatores(ativo) -> dict:
	frappe.only_for("System Manager")
	if frappe.conf.get("demo_username"):
		frappe.throw(_("No site de demonstração o login em duas etapas fica desligado (o acesso é automático)."))
	ativo = cint(ativo)
	ss = frappe.get_single("System Settings")
	ss.enable_two_factor_auth = ativo
	if ativo:
		ss.two_factor_method = "OTP App"
		ss.otp_issuer_name = (frappe.db.get_single_value("FCRM Settings", "brand_name") or "CRM")[:30]
	ss.save(ignore_permissions=True)
	for papel in PAPEIS_2FA:
		if frappe.db.exists("Role", papel):
			frappe.db.set_value("Role", papel, "two_factor_auth", ativo)
	return {"ok": True, "dois_fatores": ativo}


# ------------------------------------------------------------------ exportar tudo

TABELAS = [
	("Atendimentos", "CRM Lead", ["name", "lead_name", "email", "mobile_no", "status", "source", "indicado_por", "servico", "cidade_estado", "creation"]),
	("Casos", "CRM Deal", ["name", "lead_name", "organization", "email", "mobile_no", "status", "servico", "deal_value", "percentual_exito",
						   "numero_processo", "tribunal", "indicado_por", "encerrado", "encerrado_em", "creation"]),
	("Honorários", "CRM Honorario", ["name", "deal", "tipo_honorario", "servico", "valor", "status", "data_vencimento", "data_pagamento", "forma_pagamento", "observacoes"]),
	("Despesas", "CRM Despesa", ["name", "deal", "descricao", "categoria", "valor", "status", "data_vencimento", "data_pagamento", "reembolsavel", "reembolsado"]),
	("Recebimentos", "CRM Recebimento", ["name", "deal", "data", "origem", "valor_bruto", "percentual_honorarios", "valor_honorarios",
										 "despesas_reembolsadas", "valor_repasse", "repasse_status", "repasse_data"]),
	("Partes", "CRM Parte", ["name", "deal", "tipo", "nome", "documento"]),
	("Andamentos", "CRM Andamento", ["name", "deal", "data", "texto", "visivel_cliente"]),
	("Tarefas", "CRM Task", ["name", "title", "status", "priority", "assigned_to", "due_date", "reference_doctype", "reference_docname"]),
	("Horas", "CRM Horas", ["name", "deal", "usuario", "data", "horas", "frente", "observacoes"]),
	("Contatos", "Contact", ["name", "first_name", "last_name", "email_id", "mobile_no"]),
]


def _planilha(abas: list[tuple[str, list[str], list[list]]]) -> str:
	from openpyxl import Workbook
	from openpyxl.styles import Font

	wb = Workbook()
	wb.remove(wb.active)
	for titulo, cab, linhas in abas:
		ws = wb.create_sheet(titulo[:31])
		ws.append(cab)
		for c in ws[1]:
			c.font = Font(bold=True)
		for l in linhas:
			ws.append([v if isinstance(v, (int, float)) or v is None else str(v) for v in l])
	buf = io.BytesIO()
	wb.save(buf)
	return base64.b64encode(buf.getvalue()).decode()


@frappe.whitelist()
def exportar_tudo() -> dict:
	"""Planilha com todos os dados do escritório (uma aba por tipo)."""
	_gestor()
	abas = []
	for titulo, dt, campos in TABELAS:
		if not frappe.db.exists("DocType", dt):
			continue
		meta = frappe.get_meta(dt)
		campos = [c for c in campos if c == "name" or c == "creation" or meta.has_field(c)]
		linhas = frappe.get_all(dt, fields=campos, limit=50000, order_by="creation asc")
		abas.append((titulo, campos, [[r.get(c) for c in campos] for r in linhas]))
	marca = frappe.db.get_single_value("FCRM Settings", "brand_name") or "CRM"
	return {"arquivo": _planilha(abas), "nome": f"Dados do escritório - {marca} - {nowdate()}.xlsx"}


# ------------------------------------------------------------------ LGPD: titular

@frappe.whitelist()
def buscar_titular(q: str) -> list[dict]:
	_gestor()
	q = (q or "").strip()
	if len(q) < 3:
		return []
	like = f"%{q}%"
	out = []
	for l in frappe.get_all("CRM Lead", or_filters={"lead_name": ["like", like], "email": ["like", like], "mobile_no": ["like", like]},
							fields=["name", "lead_name", "email", "mobile_no"], limit=10):
		out.append({"tipo": "CRM Lead", "name": l.name, "rotulo": l.lead_name, "detalhe": l.email or l.mobile_no or "", "papel": _("Atendimento")})
	for d in frappe.get_all("CRM Deal", or_filters={"lead_name": ["like", like], "email": ["like", like], "mobile_no": ["like", like], "organization": ["like", like]},
							fields=["name", "lead_name", "organization", "email", "mobile_no"], limit=10):
		out.append({"tipo": "CRM Deal", "name": d.name, "rotulo": d.lead_name or d.organization or d.name, "detalhe": d.email or d.mobile_no or "", "papel": _("Caso")})
	return out


def _registros_do_titular(tipo: str, name: str) -> list[tuple[str, list[str], list[list]]]:
	doc = frappe.get_doc(tipo, name)
	campos = [f.fieldname for f in doc.meta.fields if f.fieldtype not in ("Section Break", "Column Break", "Tab Break", "HTML", "Table", "Button")
			  and f.fieldname not in ("portal_token", "proposta_token", "ficha_transcricao")]
	abas = [("Cadastro", ["campo", "valor"], [[doc.meta.get_label(c) or c, doc.get(c)] for c in campos if doc.get(c) not in (None, "", 0)])]
	if tipo == "CRM Deal":
		for titulo, dt, cs in TABELAS[2:8]:
			if frappe.db.exists("DocType", dt) and frappe.get_meta(dt).has_field("deal"):
				meta = frappe.get_meta(dt)
				cs = [c for c in cs if c == "name" or meta.has_field(c)]
				abas.append((titulo, cs, [[r.get(c) for c in cs] for r in frappe.get_all(dt, filters={"deal": name}, fields=cs)]))
	return abas


@frappe.whitelist()
def exportar_titular(tipo: str, name: str) -> dict:
	"""Tudo o que o escritório guarda sobre uma pessoa (pedido de acesso do titular, LGPD art. 18)."""
	_gestor()
	if tipo not in ("CRM Lead", "CRM Deal"):
		frappe.throw(_("Tipo inválido."))
	rot = frappe.db.get_value(tipo, name, "lead_name") or name
	return {"arquivo": _planilha(_registros_do_titular(tipo, name)), "nome": f"Dados do titular - {rot}.xlsx"}


ANON = "Titular anonimizado"


@frappe.whitelist()
def anonimizar_titular(tipo: str, name: str, confirmacao: str = "") -> dict:
	"""Apaga os dados pessoais a pedido do titular. O financeiro (honorários, recebimentos, despesas)
	continua, porque o escritório tem obrigação legal de guardar esses registros."""
	frappe.only_for(("System Manager", "Sales Manager"))
	if tipo not in ("CRM Lead", "CRM Deal"):
		frappe.throw(_("Tipo inválido."))
	if (confirmacao or "").strip().upper() != "ANONIMIZAR":
		frappe.throw(_("Digite ANONIMIZAR para confirmar."))
	doc = frappe.get_doc(tipo, name)
	limpar = {"first_name": ANON, "last_name": "", "lead_name": ANON, "email": None, "mobile_no": None, "phone": None,
			  "ficha": None, "ficha_transcricao": None, "dados_cliente": None, "instagram_sender_id": None, "instagram_username": None,
			  "instagram_photo": None, "website": None, "linkedin": None, "facebook": None, "twitter": None, "cidade_estado": None,
			  "consentimento_lgpd": 0, "consentimento_em": None, "portal_token": None}
	valores = {k: v for k, v in limpar.items() if doc.meta.has_field(k)}
	frappe.db.set_value(tipo, name, valores, update_modified=False)
	contatos = []
	if tipo == "CRM Deal":
		contatos = [c.contact for c in doc.get("contacts") or [] if c.contact]
		for p in frappe.get_all("CRM Parte", filters={"deal": name}, pluck="name"):
			frappe.db.set_value("CRM Parte", p, {"documento": None}, update_modified=False)
	for c in contatos:
		if frappe.db.exists("Contact", c):
			frappe.db.delete("Contact Email", {"parent": c})
			frappe.db.delete("Contact Phone", {"parent": c})
			frappe.db.set_value("Contact", c, {"first_name": ANON, "last_name": "", "email_id": None, "mobile_no": None, "phone": None}, update_modified=False)
	frappe.db.delete("Communication", {"reference_doctype": tipo, "reference_name": name})
	frappe.db.delete("Comment", {"reference_doctype": tipo, "reference_name": name, "comment_type": ["in", ["Comment", "Info", "Edit"]]})
	for f in frappe.get_all("File", filters={"attached_to_doctype": tipo, "attached_to_name": name}, pluck="name"):
		frappe.delete_doc("File", f, force=True, ignore_permissions=True)
	if tipo == "CRM Lead" and frappe.db.exists("DocType", "CRM Instagram Message"):
		frappe.db.delete("CRM Instagram Message", {"lead": name})
	frappe.get_doc({"doctype": "Comment", "comment_type": "Info", "reference_doctype": tipo, "reference_name": name,
					"content": _("Dados pessoais anonimizados a pedido do titular (LGPD) em {0} por {1}. Registros financeiros mantidos por obrigação legal.")
					.format(frappe.utils.format_datetime(now_datetime(), "dd/MM/yyyy HH:mm"), frappe.session.user)}).insert(ignore_permissions=True)
	return {"ok": True}


@frappe.whitelist()
def quem_acessou(deal: str) -> list[dict]:
	"""Quem abriu o caso nos últimos 90 dias (registro de acesso)."""
	_gestor()
	logs = frappe.get_all("View Log", filters={"reference_doctype": "CRM Deal", "reference_name": deal, "creation": [">=", add_days(nowdate(), -90)]},
						  fields=["viewed_by", "creation"], order_by="creation desc", limit=500)
	por_pessoa = {}
	for l in logs:
		p = por_pessoa.setdefault(l.viewed_by, {"usuario": l.viewed_by, "vezes": 0, "ultimo": str(l.creation)[:16]})
		p["vezes"] += 1
	for p in por_pessoa.values():
		p["nome"] = frappe.db.get_value("User", p["usuario"], "full_name") or p["usuario"]
	return list(por_pessoa.values())


# ------------------------------------------------------------------ carga da equipe

@frappe.whitelist()
def carga_equipe() -> list[dict]:
	_gestor()
	won = frappe.get_all("CRM Deal Status", filters={"type": "Won"}, pluck="name")
	tem_enc = frappe.get_meta("CRM Deal").has_field("encerrado")
	usuarios = frappe.get_all("Has Role", filters={"role": ["in", ["Sales User", "Sales Manager"]], "parenttype": "User"}, pluck="parent", distinct=True)
	inicio_mes = get_first_day(nowdate())
	out = []
	for u in sorted(set(usuarios)):
		info = frappe.db.get_value("User", u, ["full_name", "enabled"], as_dict=True)
		if not info or not info.enabled or u in ("Administrator", "Guest"):
			continue
		filtros = {"deal_owner": u, "status": ["in", won]}
		if tem_enc:
			filtros["encerrado"] = 0
		abertas = frappe.db.count("CRM Task", {"assigned_to": u, "status": ["in", ["Backlog", "Todo", "In Progress"]]})
		atrasadas = frappe.db.count("CRM Task", {"assigned_to": u, "status": ["in", ["Backlog", "Todo", "In Progress"]], "due_date": ["<", nowdate()]})
		horas = flt(frappe.db.sql("select coalesce(sum(horas),0) from `tabCRM Horas` where usuario=%s and data>=%s", (u, inicio_mes))[0][0]) \
			if frappe.db.exists("DocType", "CRM Horas") else 0
		out.append({"usuario": u, "nome": info.full_name or u, "casos_ativos": frappe.db.count("CRM Deal", filtros),
					"tarefas_abertas": abertas, "tarefas_atrasadas": atrasadas, "horas_mes": round(horas, 1)})
	return sorted(out, key=lambda x: (x["tarefas_atrasadas"], x["tarefas_abertas"]), reverse=True)

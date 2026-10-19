# Copyright (c) 2026, Stratcompany and contributors
# For license information, please see license.txt
#
# Importar planilha: o escritório traz o que já usa hoje (Excel ou CSV) sem redigitar nada.
# Quatro tipos: contatos, clientes e casos, honorários (financeiro) e despesas. O CRM reconhece
# sozinho as colunas pelo nome ("Cliente", "Valor", "Vencimento"...), mostra uma prévia com os
# problemas linha a linha e só grava quando a pessoa confirma. Nada é apagado nem sobrescrito:
# linha repetida (mesmo e-mail ou telefone de um contato que já existe) é ignorada.

import base64
import csv
import io
import json
import re
import unicodedata
from datetime import date, datetime

import frappe
from frappe import _
from frappe.utils import add_days, cint, flt, getdate, nowdate

MAX_LINHAS = 2000
MAX_BYTES = 5 * 1024 * 1024
PREVIA = 10

# campo: (rótulo, nomes de coluna aceitos, obrigatório, tipo)
TIPOS = {
	"contatos": {
		"titulo": "Contatos",
		"descricao": "Pessoas que procuraram o escritório e ainda não são clientes.",
		"campos": {
			"nome": ("Nome", ["nome", "nome completo", "contato", "cliente", "nome do cliente", "pessoa"], True, "texto"),
			"telefone": ("Telefone", ["telefone", "celular", "whatsapp", "zap", "fone", "tel", "telefone celular"], False, "telefone"),
			"email": ("E-mail", ["email", "e mail", "correio eletronico"], False, "email"),
			"area": ("Área", ["area", "area do direito", "servico", "assunto", "tipo de caso", "materia"], False, "texto"),
			"origem": ("Origem", ["origem", "como conheceu", "indicacao", "indicado por", "fonte", "canal"], False, "texto"),
			"cidade": ("Cidade", ["cidade", "cidade uf", "cidade estado", "municipio", "local"], False, "texto"),
			"observacoes": ("Observações", ["observacoes", "observacao", "obs", "anotacoes", "notas", "detalhes", "resumo"], False, "texto"),
		},
	},
	"casos": {
		"titulo": "Clientes e casos",
		"descricao": "Clientes do escritório e os casos de cada um.",
		"campos": {
			"nome": ("Cliente", ["cliente", "nome", "nome do cliente", "nome completo", "contratante"], True, "texto"),
			"telefone": ("Telefone", ["telefone", "celular", "whatsapp", "zap", "fone", "tel"], False, "telefone"),
			"email": ("E-mail", ["email", "e mail"], False, "email"),
			"empresa": ("Empresa", ["empresa", "razao social", "pessoa juridica", "cnpj nome"], False, "texto"),
			"area": ("Área", ["area", "area do direito", "servico", "assunto", "tipo de caso", "materia", "acao"], False, "texto"),
			"valor": ("Valor do contrato", ["valor", "valor do contrato", "honorarios", "valor total", "valor contratado"], False, "valor"),
			"situacao": ("Situação", ["situacao", "status", "andamento", "fase", "etapa"], False, "texto"),
			"observacoes": ("Observações", ["observacoes", "observacao", "obs", "anotacoes", "notas", "detalhes", "descricao do caso"], False, "texto"),
		},
	},
	"honorarios": {
		"titulo": "Honorários (financeiro)",
		"descricao": "Parcelas e pagamentos de honorários: pagos, a receber e atrasados.",
		"campos": {
			"cliente": ("Cliente", ["cliente", "nome", "nome do cliente", "contratante", "pagador"], True, "texto"),
			"valor": ("Valor", ["valor", "valor da parcela", "honorario", "honorarios", "quantia", "total"], True, "valor"),
			"vencimento": ("Vencimento", ["vencimento", "data de vencimento", "vence em", "data", "previsao"], False, "data"),
			"pago_em": ("Pago em", ["pago em", "data de pagamento", "pagamento", "recebido em", "data do pagamento", "quitado em"], False, "data"),
			"situacao": ("Situação", ["situacao", "status", "pago", "pago?", "quitado"], False, "texto"),
			"tipo": ("Tipo", ["tipo", "tipo de honorario", "modalidade"], False, "texto"),
			"forma": ("Forma de pagamento", ["forma", "forma de pagamento", "meio de pagamento", "meio"], False, "texto"),
			"area": ("Área", ["area", "area do direito", "servico", "assunto", "materia"], False, "texto"),
			"observacoes": ("Descrição", ["descricao", "referente", "referencia", "observacoes", "obs", "parcela", "historico"], False, "texto"),
		},
	},
	"despesas": {
		"titulo": "Despesas e custas",
		"descricao": "Custas, despesas do escritório e o que foi adiantado por cliente.",
		"campos": {
			"descricao": ("Descrição", ["descricao", "despesa", "item", "historico", "referente"], True, "texto"),
			"valor": ("Valor", ["valor", "quantia", "total", "custo"], True, "valor"),
			"vencimento": ("Vencimento", ["vencimento", "data de vencimento", "data", "vence em"], False, "data"),
			"pago_em": ("Pago em", ["pago em", "data de pagamento", "pagamento", "quitado em"], False, "data"),
			"situacao": ("Situação", ["situacao", "status", "pago", "pago?"], False, "texto"),
			"categoria": ("Categoria", ["categoria", "tipo", "grupo", "natureza"], False, "texto"),
			"fornecedor": ("Fornecedor", ["fornecedor", "favorecido", "pago a", "beneficiario"], False, "texto"),
			"cliente": ("Cliente (se for custa de um caso)", ["cliente", "nome do cliente", "caso"], False, "texto"),
		},
	},
}

EXEMPLOS = {
	"contatos": [
		["Maria Souza", "(87) 99999-1111", "maria@email.com", "Trabalhista", "Indicação", "Petrolina/PE", "Demitida sem receber a rescisão"],
		["João Lima", "(87) 98888-2222", "", "Previdenciário", "Instagram", "Juazeiro/BA", "Aposentadoria negada"],
	],
	"casos": [
		["Ana Pereira", "(87) 99777-3333", "ana@email.com", "", "Família e Sucessões", "4.500,00", "Ativo", "Divórcio consensual"],
		["Carlos Melo", "(87) 99666-4444", "", "Mercado Central Ltda", "Consumidor", "1.200,00", "Em negociação", "Assessoria mensal"],
	],
	"honorarios": [
		["Ana Pereira", "1.500,00", "10/09/2026", "10/09/2026", "Pago", "Fixo", "PIX", "Família e Sucessões", "Parcela 1 de 3"],
		["Ana Pereira", "1.500,00", "10/10/2026", "", "Pendente", "Fixo", "PIX", "Família e Sucessões", "Parcela 2 de 3"],
	],
	"despesas": [
		["Custas iniciais", "350,00", "05/09/2026", "05/09/2026", "Pago", "Custas", "Tribunal", "Ana Pereira"],
		["Aluguel da sala", "2.200,00", "05/10/2026", "", "Pendente", "Contas fixas", "Imobiliária", ""],
	],
}


# ------------------------------------------------------------------ utilidades

def _managers_only():
	frappe.only_for(("System Manager", "Sales Manager"))


def _norm(text) -> str:
	t = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode().lower()
	return re.sub(r"[^a-z0-9]+", " ", t).strip()


def _tipo(tipo: str) -> dict:
	if tipo not in TIPOS:
		frappe.throw(_("Tipo de planilha desconhecido."))
	return TIPOS[tipo]


def _ler_arquivo(arquivo: str, nome_arquivo: str) -> list[list]:
	"""Lê .xlsx ou .csv (base64) e devolve as linhas como listas de valores."""
	try:
		data = base64.b64decode(arquivo.split(",", 1)[-1])
	except Exception:
		frappe.throw(_("Não foi possível ler o arquivo."))
	if len(data) > MAX_BYTES:
		frappe.throw(_("Arquivo maior que 5 MB. Divida a planilha em partes."))
	ext = (nome_arquivo or "").rsplit(".", 1)[-1].lower()
	if ext in ("xlsx", "xlsm"):
		from openpyxl import load_workbook

		try:
			wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
		except Exception:
			frappe.throw(_("A planilha parece estar corrompida. Salve de novo como .xlsx e tente outra vez."))
		ws = wb.worksheets[0]
		rows = [list(r) for r in ws.iter_rows(values_only=True)]
		wb.close()
	elif ext == "csv":
		text = None
		for enc in ("utf-8-sig", "latin-1"):
			try:
				text = data.decode(enc)
				break
			except UnicodeDecodeError:
				continue
		sample = text[:4096]
		delim = ";" if sample.count(";") > sample.count(",") else ","
		rows = list(csv.reader(io.StringIO(text), delimiter=delim))
	elif ext == "xls":
		frappe.throw(_("Esse é o formato antigo do Excel (.xls). Abra no Excel e salve como .xlsx."))
	else:
		frappe.throw(_("Envie uma planilha .xlsx (Excel) ou .csv."))
	return [r for r in rows if any(str(c).strip() for c in r if c is not None)]


def _cabecalho(rows: list[list], campos: dict) -> int:
	"""Acha a linha do cabeçalho: a primeira (entre as 10 primeiras) com mais colunas reconhecidas."""
	best, best_hits = 0, -1
	for i, row in enumerate(rows[:10]):
		hits = sum(1 for c in row if _campo_da_coluna(c, campos))
		if hits > best_hits:
			best, best_hits = i, hits
	return best


def _campo_da_coluna(coluna, campos: dict) -> str | None:
	n = _norm(coluna)
	if not n:
		return None
	for key, (label, sinonimos, _r, _t) in campos.items():
		if n == _norm(label) or n in sinonimos:
			return key
	for key, (label, sinonimos, _r, _t) in campos.items():
		if any(len(s) > 3 and (n.startswith(s) or s in n.split(" ") or n in s.split(" ") and len(n) > 3) for s in sinonimos):
			return key
	return None


def _mapear(cabecalho: list, campos: dict) -> dict:
	usado, mapa = set(), {}
	for idx, col in enumerate(cabecalho):
		key = _campo_da_coluna(col, campos)
		if key and key not in mapa and idx not in usado:
			mapa[key] = idx
			usado.add(idx)
	return {k: mapa.get(k) for k in campos}


def _valor(v):
	if v is None or v == "":
		return None
	if isinstance(v, (int, float)):
		return round(float(v), 2)
	s = re.sub(r"[^0-9,.\-]", "", str(v))
	if not s:
		return None
	if "," in s:
		s = s.replace(".", "").replace(",", ".")
	elif s.count(".") > 1:
		s = s.replace(".", "")
	try:
		return round(float(s), 2)
	except ValueError:
		return None


def _data(v):
	if v is None or v == "":
		return None
	if isinstance(v, datetime):
		return v.date()
	if isinstance(v, date):
		return v
	if isinstance(v, (int, float)) and 20000 < v < 80000:  # número de série do Excel
		return add_days("1899-12-30", int(v))
	s = str(v).strip()
	m = re.match(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})", s)
	if m:
		d, mth, y = (int(x) for x in m.groups())
		y = y + 2000 if y < 100 else y
		try:
			return date(y, mth, d)
		except ValueError:
			return None
	m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s)
	if m:
		try:
			return date(*(int(x) for x in m.groups()))
		except ValueError:
			return None
	return None


def _texto(v) -> str:
	if v is None:
		return ""
	if isinstance(v, float) and v.is_integer():
		v = int(v)
	return str(v).strip()


def _email(v) -> str:
	s = _texto(v).lower()
	return s if re.match(r"^[^@\s]+@[^@\s]+\.[a-z]{2,}$", s) else ""


def _telefone(v) -> str:
	s = _texto(v)
	return s if len(re.sub(r"\D", "", s)) >= 8 else ""


def _linha(row: list, mapa: dict, campos: dict) -> tuple[dict, list[str]]:
	"""Converte uma linha usando o mapeamento e devolve (valores, problemas)."""
	vals, erros = {}, []
	for key, (label, _s, obrigatorio, tipo) in campos.items():
		idx = mapa.get(key)
		bruto = row[idx] if idx is not None and idx < len(row) else None
		if tipo == "valor":
			v = _valor(bruto)
			if bruto not in (None, "") and v is None:
				erros.append(_("{0}: valor não reconhecido").format(label))
		elif tipo == "data":
			v = _data(bruto)
			if bruto not in (None, "") and v is None:
				erros.append(_("{0}: data não reconhecida (use dd/mm/aaaa)").format(label))
		elif tipo == "email":
			v = _email(bruto)
			if _texto(bruto) and not v:
				erros.append(_("{0}: e-mail inválido (a linha entra sem e-mail)").format(label))
		elif tipo == "telefone":
			v = _telefone(bruto)
		else:
			v = _texto(bruto)
		if obrigatorio and (v is None or v == ""):
			erros.append(_("{0} está em branco").format(label))
		vals[key] = v.isoformat() if isinstance(v, date) else v
	return vals, erros


def _bloqueia(erros: list[str]) -> bool:
	"""Erros que impedem a linha de entrar (os de e-mail inválido não impedem)."""
	return any("inválido (a linha entra" not in e for e in erros)


def _preparar(tipo: str, arquivo: str, nome_arquivo: str, mapeamento=None):
	cfg = _tipo(tipo)
	campos = cfg["campos"]
	rows = _ler_arquivo(arquivo, nome_arquivo)
	if not rows:
		frappe.throw(_("A planilha está vazia."))
	h = _cabecalho(rows, campos)
	cabecalho = [_texto(c) or _("Coluna {0}").format(i + 1) for i, c in enumerate(rows[h])]
	dados = rows[h + 1 :]
	if len(dados) > MAX_LINHAS:
		frappe.throw(_("A planilha tem mais de {0} linhas. Divida em partes.").format(MAX_LINHAS))
	if mapeamento:
		mapa = frappe.parse_json(mapeamento) if isinstance(mapeamento, str) else mapeamento
		mapa = {k: (cint(v) if v not in (None, "", -1, "-1") else None) for k, v in mapa.items() if k in campos}
		for k in campos:
			mapa.setdefault(k, None)
	else:
		mapa = _mapear(cabecalho, campos)
	return cfg, campos, cabecalho, dados, h, mapa


# ------------------------------------------------------------------ API

@frappe.whitelist()
def tipos() -> list[dict]:
	_managers_only()
	return [
		{
			"tipo": k,
			"titulo": v["titulo"],
			"descricao": v["descricao"],
			"campos": [{"campo": c, "rotulo": l, "obrigatorio": o} for c, (l, _s, o, _t) in v["campos"].items()],
		}
		for k, v in TIPOS.items()
	]


@frappe.whitelist()
def baixar_modelo(tipo: str) -> dict:
	"""Planilha modelo (.xlsx) com as colunas e duas linhas de exemplo."""
	_managers_only()
	cfg = _tipo(tipo)
	from openpyxl import Workbook
	from openpyxl.styles import Font, PatternFill

	wb = Workbook()
	ws = wb.active
	ws.title = cfg["titulo"][:31]
	ws.append([l for (l, _s, _o, _t) in cfg["campos"].values()])
	for cell in ws[1]:
		cell.font = Font(bold=True, color="FFFFFF")
		cell.fill = PatternFill("solid", fgColor="042D3C")
	for ex in EXEMPLOS[tipo]:
		ws.append(ex)
	for col in ws.columns:
		ws.column_dimensions[col[0].column_letter].width = max(14, min(40, max(len(str(c.value or "")) for c in col) + 2))
	buf = io.BytesIO()
	wb.save(buf)
	return {"arquivo": base64.b64encode(buf.getvalue()).decode(), "nome": f"Modelo - {cfg['titulo']}.xlsx"}


@frappe.whitelist()
def analisar(tipo: str, arquivo: str, nome_arquivo: str, mapeamento=None) -> dict:
	"""Lê a planilha, reconhece as colunas e devolve a prévia (nada é gravado)."""
	_managers_only()
	cfg, campos, cabecalho, dados, h, mapa = _preparar(tipo, arquivo, nome_arquivo, mapeamento)
	previa, validas, com_erro = [], 0, 0
	for i, row in enumerate(dados):
		vals, erros = _linha(row, mapa, campos)
		if _bloqueia(erros):
			com_erro += 1
		else:
			validas += 1
		if len(previa) < PREVIA or (erros and len([p for p in previa if p["erros"]]) < 30):
			previa.append({"linha": h + i + 2, "valores": vals, "erros": erros})
	faltando = [campos[k][0] for k, (_l, _s, o, _t) in campos.items() if o and mapa.get(k) is None]
	return {
		"colunas": cabecalho,
		"mapeamento": mapa,
		"previa": previa,
		"total": len(dados),
		"validas": validas,
		"com_erro": com_erro,
		"obrigatorios_faltando": faltando,
	}


@frappe.whitelist()
def importar(tipo: str, arquivo: str, nome_arquivo: str, mapeamento=None) -> dict:
	"""Grava as linhas válidas. Linha com problema é pulada e aparece no relatório."""
	_managers_only()
	cfg, campos, cabecalho, dados, h, mapa = _preparar(tipo, arquivo, nome_arquivo, mapeamento)
	faltando = [campos[k][0] for k, (_l, _s, o, _t) in campos.items() if o and mapa.get(k) is None]
	if faltando:
		frappe.throw(_("Escolha a coluna de: {0}").format(", ".join(faltando)))
	criar = {"contatos": _criar_contato, "casos": _criar_caso, "honorarios": _criar_honorario, "despesas": _criar_despesa}[tipo]
	ctx = {"casos": {}}
	criados = ignorados = 0
	problemas = []
	for i, row in enumerate(dados):
		n = h + i + 2
		vals, erros = _linha(row, mapa, campos)
		if _bloqueia(erros):
			problemas.append({"linha": n, "motivo": "; ".join(erros)})
			continue
		sp = f"imp_{i}"
		frappe.db.savepoint(sp)
		try:
			if criar(vals, ctx):
				criados += 1
			else:
				ignorados += 1
		except Exception as e:
			frappe.db.rollback(save_point=sp)
			problemas.append({"linha": n, "motivo": _("não foi possível gravar: {0}").format(str(e)[:160])})
	frappe.db.commit()
	return {"criados": criados, "ignorados": ignorados, "com_erro": len(problemas), "problemas": problemas[:100]}


# ------------------------------------------------------------------ gravação

def _dividir_nome(nome: str) -> tuple[str, str]:
	partes = nome.split()
	return (partes[0], " ".join(partes[1:])) if len(partes) > 1 else (nome, "")


def _parecido(alvo: str, opcoes: list[str]) -> str | None:
	return next((o for o in opcoes if _norm(o) == alvo), None) or next(
		(o for o in opcoes if alvo and len(alvo) > 3 and (alvo in _norm(o) or _norm(o) in alvo)), None
	)


def _area(texto: str) -> str | None:
	"""Área da planilha -> Serviço do CRM. Usa a área do perfil do escritório quando bate
	("Família" -> "Família e Sucessões"); se não existir, cria o serviço para não perder a informação."""
	if not texto or not _norm(texto):
		return None
	alvo = _norm(texto)
	opcoes = [o for o in (frappe.get_meta("CRM Honorario").get_field("servico").options or "").split("\n") if o]
	existentes = frappe.get_all("CRM Servico", pluck="name")
	escolhida = _parecido(alvo, opcoes) or _parecido(alvo, existentes) or texto.strip()[:60]
	if not frappe.db.exists("CRM Servico", escolhida):
		frappe.get_doc({"doctype": "CRM Servico", "nome": escolhida, "ativo": 1}).insert(ignore_permissions=True)
	return escolhida


def _area_financeiro(area: str | None) -> str | None:
	"""O campo de área do honorário só aceita as áreas do perfil; fora delas fica em branco."""
	opcoes = (frappe.get_meta("CRM Honorario").get_field("servico").options or "").split("\n")
	return area if area in opcoes else None


ORIGENS = {
	"indicacao": "Reference", "indicado": "Reference", "indicada": "Reference", "cliente": "Existing Customer",
	"instagram": "Facebook", "facebook": "Facebook", "anuncio": "Advertisement", "google": "Website",
	"site": "Website", "website": "Website", "presencial": "Walk In", "balcao": "Walk In", "email": "Email",
}


def _origem(texto: str) -> str | None:
	n = _norm(texto)
	if not n:
		return None
	for chave, fonte in ORIGENS.items():
		if chave in n.split(" ") or n.startswith(chave):
			return fonte
	nome = texto.strip()[:60]
	if not frappe.db.exists("CRM Lead Source", nome):
		frappe.get_doc({"doctype": "CRM Lead Source", "source_name": nome}).insert(ignore_permissions=True)
	return nome


def _ja_existe_contato(email: str, telefone: str) -> bool:
	digitos = re.sub(r"\D", "", telefone or "")[-8:]
	if email and (frappe.db.exists("CRM Lead", {"email": email}) or frappe.db.exists("CRM Deal", {"email": email})):
		return True
	if digitos:
		for dt in ("CRM Lead", "CRM Deal"):
			if frappe.db.sql(
				f"select name from `tab{dt}` where replace(replace(replace(replace(replace(coalesce(mobile_no,''),' ',''),'-',''),'(',''),')',''),'+','') like %s limit 1",
				(f"%{digitos}",),
			):
				return True
	return False


def _criar_contato(v: dict, ctx) -> bool:
	if _ja_existe_contato(v.get("email"), v.get("telefone")):
		return False
	first, last = _dividir_nome(v["nome"])
	nota = v.get("observacoes") or ""
	frappe.get_doc(
		{
			"doctype": "CRM Lead",
			"first_name": first,
			"last_name": last,
			"email": v.get("email") or None,
			"mobile_no": v.get("telefone") or None,
			"status": "New",
			"source": _origem(v.get("origem")),
			"servico": _area(v.get("area")),
			"cidade_estado": v.get("cidade") or None,
			"ficha": json.dumps({"dor_objetivo": nota}, ensure_ascii=False) if nota else None,
		}
	).insert(ignore_permissions=True)
	return True


SITUACAO_CASO = {
	"ativo": "Won", "andamento": "Won", "em andamento": "Won", "fechado": "Won", "ganho": "Won", "contratado": "Won",
	"cliente": "Won", "encerrado": "Won", "concluido": "Won", "finalizado": "Won", "arquivado": "Won",
	"perdido": "Lost", "desistiu": "Lost", "nao fechou": "Lost", "recusou": "Lost",
	"proposta": "Proposal/Quotation", "proposta enviada": "Proposal/Quotation", "negociacao": "Negotiation",
	"em negociacao": "Negotiation",
}


def _status_caso(texto: str) -> str:
	n = _norm(texto)
	if n in SITUACAO_CASO:
		return SITUACAO_CASO[n]
	for chave, st in SITUACAO_CASO.items():
		if chave in n:
			return st
	return "Qualification"


def _motivo_importado() -> str:
	nome = "Importado da planilha"
	if not frappe.db.exists("CRM Lost Reason", nome):
		frappe.get_doc({"doctype": "CRM Lost Reason", "lost_reason": nome}).insert(ignore_permissions=True)
	return nome


def _novo_caso(nome: str, telefone="", email="", empresa="", area=None, valor=None, status="Won", nota="") -> str:
	"""Cria o caso já com a situação certa, sem disparar a cobrança automática de negócio ganho
	(o histórico vem da planilha; parcelas antigas entram pelo import de honorários)."""
	first, last = _dividir_nome(nome)
	doc = frappe.get_doc(
		{
			"doctype": "CRM Deal",
			"status": status,
			"first_name": first,
			"last_name": last,
			"email": email or None,
			"mobile_no": telefone or None,
			"servico": area,
			"deal_value": valor or None,
			"lead_name": nome,
			"ficha": json.dumps({"dor_objetivo": nota}, ensure_ascii=False) if nota else None,
		}
	)
	if empresa:
		if not frappe.db.exists("CRM Organization", empresa):
			frappe.get_doc({"doctype": "CRM Organization", "organization_name": empresa}).insert(ignore_permissions=True)
		doc.organization = empresa
	if status == "Lost":
		doc.lost_reason = _motivo_importado()
	if email or telefone:
		c = frappe.get_doc({"doctype": "Contact", "first_name": first, "last_name": last})
		if email:
			c.append("email_ids", {"email_id": email, "is_primary": 1})
		if telefone:
			c.append("phone_nos", {"phone": telefone, "is_primary_mobile_no": 1})
		c.insert(ignore_permissions=True)
		doc.append("contacts", {"contact": c.name, "is_primary": 1})
	doc.flags.sem_cobranca_automatica = True
	doc.insert(ignore_permissions=True)
	return doc.name


def _achar_caso(nome: str, ctx) -> str | None:
	chave = _norm(nome)
	if chave in ctx["casos"]:
		return ctx["casos"][chave]
	found = None
	for d in frappe.get_all(
		"CRM Deal", fields=["name", "lead_name", "first_name", "last_name", "organization"], order_by="creation desc", limit=5000
	):
		nomes = [d.lead_name, f"{d.first_name or ''} {d.last_name or ''}", d.organization]
		if any(_norm(x) == chave for x in nomes if x):
			found = d.name
			break
	ctx["casos"][chave] = found
	return found


def _criar_caso(v: dict, ctx) -> bool:
	if _ja_existe_contato(v.get("email"), v.get("telefone")) or _achar_caso(v["nome"], ctx):
		return False
	nome = _novo_caso(
		v["nome"], v.get("telefone"), v.get("email"), v.get("empresa"), _area(v.get("area")),
		v.get("valor"), _status_caso(v.get("situacao")), v.get("observacoes") or "",
	)
	ctx["casos"][_norm(v["nome"])] = nome
	return True


def _situacao_pagamento(texto: str, pago_em, vencimento) -> str:
	n = _norm(texto)
	if n in ("pago", "paga", "sim", "s", "quitado", "quitada", "recebido", "recebida", "ok") or pago_em:
		return "Pago"
	if n in ("atrasado", "atrasada", "vencido", "vencida", "em atraso"):
		return "Atrasado"
	if vencimento and getdate(vencimento) < getdate(nowdate()) and n not in ("pendente", "aberto", "a receber", "nao", "n"):
		return "Atrasado"
	return "Pendente"


def _tipo_honorario(texto: str) -> str:
	n = _norm(texto)
	if any(k in n for k in ("exito", "sucesso", "resultado")):
		return "Êxito"
	if any(k in n for k in ("mensal", "partido", "recorrente", "consultivo", "mensalidade")):
		return "Consultivo Mensal"
	if "ato" in n.split(" "):
		return "Por Ato"
	return "Fixo"


def _forma(texto: str) -> str | None:
	n = _norm(texto)
	for chave, forma in (("pix", "PIX"), ("boleto", "Boleto"), ("cartao", "Cartão"), ("credito", "Cartão"),
						 ("debito", "Cartão"), ("transfer", "Transferência"), ("ted", "Transferência"), ("doc", "Transferência")):
		if chave in n:
			return forma
	return None


def _criar_honorario(v: dict, ctx) -> bool:
	deal = _achar_caso(v["cliente"], ctx)
	area = _area(v.get("area"))
	if not deal:
		deal = _novo_caso(v["cliente"], area=area, status="Won")
		ctx["casos"][_norm(v["cliente"])] = deal
	vencimento = v.get("vencimento") or v.get("pago_em") or nowdate()
	situacao = _situacao_pagamento(v.get("situacao"), v.get("pago_em"), vencimento)
	frappe.get_doc(
		{
			"doctype": "CRM Honorario",
			"deal": deal,
			"status": situacao,
			"tipo_honorario": _tipo_honorario(v.get("tipo")),
			"servico": _area_financeiro(area),
			"valor": v["valor"],
			"forma_pagamento": _forma(v.get("forma")),
			"data_vencimento": vencimento,
			"data_pagamento": (v.get("pago_em") or vencimento) if situacao == "Pago" else None,
			"observacoes": (v.get("observacoes") or _("Importado da planilha"))[:500],
		}
	).insert(ignore_permissions=True)
	return True


# palavra na planilha -> categorias candidatas (usa a primeira que existir no perfil do site)
CATEGORIAS = {
	"custa": ["Custas e Taxas"], "taxa": ["Custas e Taxas"],
	"pericia": ["Perícias e Diligências", "Custas e Taxas"], "diligencia": ["Perícias e Diligências", "Custas e Taxas"],
	"cartorio": ["Cartório e Certidões", "Custas e Taxas"], "certid": ["Cartório e Certidões", "Custas e Taxas"],
	"correspondente": ["Correspondentes", "Freelancers"], "deslocamento": ["Deslocamento"], "combustivel": ["Deslocamento"],
	"uber": ["Deslocamento"], "viagem": ["Deslocamento"], "imposto": ["Impostos e Taxas"], "das": ["Impostos e Taxas"],
	"aluguel": ["Contas Fixas"], "luz": ["Contas Fixas"], "energia": ["Contas Fixas"], "agua": ["Contas Fixas"],
	"internet": ["Contas Fixas"], "telefone": ["Contas Fixas"], "fixa": ["Contas Fixas"],
	"software": ["Ferramentas e Software"], "sistema": ["Ferramentas e Software"], "assinatura": ["Ferramentas e Software"],
	"anuncio": ["Anúncios e Marketing", "Tráfego Pago"], "trafego": ["Anúncios e Marketing", "Tráfego Pago"],
	"marketing": ["Anúncios e Marketing", "Tráfego Pago"], "equipamento": ["Equipamentos"], "computador": ["Equipamentos"],
	"freelancer": ["Freelancers", "Correspondentes"],
}


def _categoria(texto: str, descricao: str) -> str:
	opcoes = [o for o in (frappe.get_meta("CRM Despesa").get_field("categoria").options or "").split("\n") if o]
	n = _norm(texto)
	exata = next((o for o in opcoes if _norm(o) == n), None)
	if exata:
		return exata
	alvo = f"{n} {_norm(descricao)}"
	for chave, candidatas in CATEGORIAS.items():
		if chave in alvo:
			achada = next((c for c in candidatas if c in opcoes), None)
			if achada:
				return achada
	return "Outros" if "Outros" in opcoes else (opcoes[-1] if opcoes else None)


def _criar_despesa(v: dict, ctx) -> bool:
	vencimento = v.get("vencimento") or v.get("pago_em") or nowdate()
	situacao = _situacao_pagamento(v.get("situacao"), v.get("pago_em"), vencimento)
	deal = _achar_caso(v["cliente"], ctx) if v.get("cliente") else None
	frappe.get_doc(
		{
			"doctype": "CRM Despesa",
			"descricao": v["descricao"][:140],
			"valor": v["valor"],
			"categoria": _categoria(v.get("categoria"), v["descricao"]),
			"fornecedor": v.get("fornecedor") or None,
			"status": situacao,
			"data_vencimento": vencimento,
			"data_pagamento": (v.get("pago_em") or vencimento) if situacao == "Pago" else None,
			"deal": deal,
			"observacoes": _("Importado da planilha") + (f" — {_('cliente')}: {v['cliente']}" if v.get("cliente") and not deal else ""),
		}
	).insert(ignore_permissions=True)
	return True
